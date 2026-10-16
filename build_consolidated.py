import json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

with open('all_months_final_clean.json', 'r', encoding='utf-8') as f:
    all_combined = json.load(f)

# Sort chronologically
month_order = {'Enero': 1, 'Febrero': 2, 'Marzo': 3, 'Abril': 4, 'Mayo': 5, 'Junio': 6}
all_combined.sort(key=lambda x: (month_order.get(x['mes'], 99), str(x['estab'])))

wb = openpyxl.Workbook()
ws_all = wb.active
ws_all.title = "Consolidado 2026"

sheets = {
    'Consolidado 2026': ws_all,
    'Enero': wb.create_sheet('Enero'),
    'Febrero': wb.create_sheet('Febrero'),
    'Marzo': wb.create_sheet('Marzo'),
    'Abril': wb.create_sheet('Abril'),
    'Mayo': wb.create_sheet('Mayo'),
    'Junio': wb.create_sheet('Junio')
}

headers = [
    "Mes / Carpeta",
    "CUIT",
    "Razón Social",
    "N° Est.",
    "Domicilio del Establecimiento",
    "Análisis de Agua (Detalle & Estado)",
    "P.A.T. - Puesta a Tierra (Resultado)",
    "Medición de Iluminación (Resultado)",
    "Otras Mediciones (Ruido, etc.)",
    "Nombre del Archivo"
]

header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

data_font = Font(name="Calibri", size=10)
top_left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
center_top_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

col_widths = {1: 14, 2: 16, 3: 34, 4: 10, 5: 35, 6: 48, 7: 35, 8: 40, 9: 38, 10: 45}

for s_name, ws in sheets.items():
    ws.append(headers)
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

# Populate sheets
for r in all_combined:
    row_data = [
        r['mes'],
        r['cuit'],
        r['rs'],
        r['estab'],
        r['domicilio'],
        r['agua'],
        r['pat'],
        r['iluminacion'],
        r['otras'],
        r['filename']
    ]
    # Add to Consolidado
    ws_all.append(row_data)
    curr_row = ws_all.max_row
    ws_all.row_dimensions[curr_row].height = 80
    for col_idx in range(1, 11):
        c = ws_all.cell(row=curr_row, column=col_idx)
        c.font = data_font
        c.border = thin_border
        c.alignment = center_top_align if col_idx in [1, 2, 4] else top_left_align
        
    # Add to specific month sheet
    ws_m = sheets.get(r['mes'])
    if ws_m:
        ws_m.append(row_data)
        m_row = ws_m.max_row
        ws_m.row_dimensions[m_row].height = 80
        for col_idx in range(1, 11):
            c = ws_m.cell(row=m_row, column=col_idx)
            c.font = data_font
            c.border = thin_border
            c.alignment = center_top_align if col_idx in [1, 2, 4] else top_left_align

excel_out = r'c:\Users\cript\OneDrive\Escritorio\Antigravity\Informes IAPSER\Relevamiento_Mediciones_Consolidado_Enero_Junio_2026.xlsx'
wb.save(excel_out)
print(f"Excel consolidated saved successfully to: {excel_out}")
