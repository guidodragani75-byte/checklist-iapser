import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Informes Abril - Mediciones"

# Headers requested:
# Carpeta Mes, CUIT, Razón Social, N° Establecimiento, Domicilio, Análisis de Agua, PAT (Puesta a Tierra), Medición de Iluminación, Otras Mediciones (Ruido, etc.)
headers = [
    "Mes / Carpeta",
    "CUIT",
    "Razón Social",
    "N° Est.",
    "Domicilio del Establecimiento",
    "Análisis de Agua (Resultado)",
    "P.A.T. - Puesta a Tierra (Resultado)",
    "Medición de Iluminación (Resultado)",
    "Otras Mediciones (Ruido, etc.)",
    "Nombre del Archivo"
]

rows = [
    [
        "Abril",
        "30-68112233-4",
        "MINISTERIO DE GOBIERNO Y TRABAJO DE LA PROVINCIA DE ENTRE RÍOS",
        "25",
        "Lavalle y 9 de Julio (Centro Cívico, parte posterior), Federal, Entre Ríos (CP 3180)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO.\n• Bacteriológico: NO APTO (recuentos elevados de bacterias aerobias mesófilas; sin coliformes fecales ni Pseudomonas).\n• Conclusión: NO APTA para consumo humano. Requiere limpieza y desinfección de tanques y nuevo análisis.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 0,89 Ω (Esquema TT, disyuntor DD).\n• Conclusión: ADECUADO / CONFORME (muy por debajo de los 40 Ω máximos reglamentarios).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• 3 puntos medidos: Mesa de entradas (278 lux vs 300 req), Relaciones laborales (226 lux vs 300 req), Oficina delegado (341 lux vs 300 req).\n• Conclusión: NO CUMPLE en sectores administrativos (distribución heterogénea; requiere optimización de luminarias y luces de emergencia).",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica.",
        "30681122334_MINISTERIO DE GOBIERNO Y TRABAJO DE LA PROV. E.R._FEDERAL_25_Abr2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-70755869-1",
        "UNIVERSIDAD AUTÓNOMA DE ENTRE RÍOS (UADER)",
        "16",
        "Tratado del Pilar 314, Diamante, Entre Ríos (CP 3105)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO (parámetros normales, sin cloro residual).\n• Bacteriológico: NO APTO (presencia de coliformes totales, coliformes fecales y aerobios mesófilos por encima de límites).\n• Conclusión: NO APTA para consumo humano. Riesgo sanitario inmediato. Se recomendó restringir consumo de inmediato y proveer agua segura alternativa.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 0,28 Ω (Jabalina tablero principal, esquema TT con disyuntor DD).\n• Conclusión: ADECUADO / CONFORME (condiciones seguras y disipación correcta).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Conclusión: Medición efectuada con instrumental específico. Se sugieren adecuaciones en distribución lumínica en puestos de trabajo.",
        "RUIDO: Relevado.\n• Conclusión: No se realizaron mediciones instrumentales específicas por no evidenciarse exposición sonora significativa.",
        "30707558691_UADER_TECNICATURA EN ACUICULTURA FCY_16_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-70755869-1",
        "UNIVERSIDAD AUTÓNOMA DE ENTRE RÍOS (UADER)",
        "32",
        "Irigoyen 1606 / 1644 (Taller Esc. Técnica N° 35), Crespo, Entre Ríos (CP 3116)",
        "SÍ SE REALIZÓ.\n• Bacteriológico: APTO (sin contaminación microbiológica).\n• Fisicoquímico: NO APTO por ARSÉNICO ELEVADO (0,04 mg/l vs valor de referencia normativo de 0,01 mg/l).\n• Conclusión: NO APTA para consumo humano prolongado por contenido de arsénico.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 1,14 Ω (Jabalina sector interno, esquema TT con DD).\n• Conclusión: ADECUADO / CONFORME (dentro de recomendaciones legales < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Medición en 3 sectores (Talleres y aulas).\n• Conclusión: Realizada; se indican ajustes en sectores de máquinas y bancos.",
        "RUIDO LABORAL (PROTOCOLO RES. SRT 85/12):\nSÍ SE REALIZÓ con decibelímetro Wintact WT85B calibrado.\n• Puesto A (Torno 1): 71,2 dBA (CUMPLE Dec. 351/79 para 8 hs).\n• Puesto B (Torno 2): 71,3 dBA (CUMPLE Dec. 351/79).\n• Puesto C: CUMPLE.\n• Puesto D (Compresores): Desconectado momentáneamente, a reubicar fuera de taller.",
        "30707558691_UNIVERSIDADAUTONOMADEENTRERIOS_32_CRESPOESCTECNICA35_abr26_INFORMETECNICO.pdf"
    ],
    [
        "Abril",
        "30-70755869-1",
        "UNIVERSIDAD AUTÓNOMA DE ENTRE RÍOS (UADER)",
        "33",
        "San Martín 1606 (Esc. Nivel Medio y Superior José de San Martín), Crespo, Entre Ríos (CP 3116)",
        "SÍ SE REALIZÓ.\n• Bacteriológico: NO APTO (contaminación microbiológica por encima de parámetros).\n• Fisicoquímico: NO APTO por ARSÉNICO ELEVADO (0,04 mg/l vs ref 0,01 mg/l).\n• Conclusión: NO APTA para consumo humano. Requiere limpieza de tanques de agua, filtros/tratamiento y nueva medición.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 1,10 Ω (Jabalina tablero general, esquema TT con DD).\n• Conclusión: ADECUADO / CONFORME (cumple normativas legales < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Medición en 3 sectores de la institución educativa.\n• Conclusión: Valores relevados; se sugiere adecuar luminarias en aulas con tubos agotados.",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica durante la visita.",
        "30707558691_UNIVERSIDADAUTONOMADEENTRERIOS_33_ESC.JJSANMARTIN_abr26_INFORMETECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "76",
        "Suipacha 844, Ciudad Autónoma de Buenos Aires (CP 1035)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO.\n• Bacteriológico: APTO (sin coliformes totales, E. coli ni Pseudomonas aeruginosa).\n• Conclusión: APTA PARA CONSUMO HUMANO conforme al Art. 982 del Código Alimentario Argentino.",
        "NO SE PUDO REALIZAR.\n• Conclusión: No fue posible efectuar la medición por imposibilidad de localizar y acceder físicamente a la jabalina/sistema PAT del edificio. Se recomendó gestión con el consorcio/personal técnico para su medición.",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Instrumental: Luxómetro Schwyz LUX-001.\n• Conclusión: CUMPLE con los valores mínimos requeridos por el Dec. 351/79 en los puestos evaluados.",
        "RUIDO: Relevado.\n• Conclusión: No se realizaron mediciones instrumentales por no resultar necesario según el tipo de actividad administrativa.",
        "30712147829_ADMINISTRADORATRIBUTARIAER_CAPITAL FEDERAL_76_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "38",
        "Moreno 1553, Crespo, Entre Ríos (CP 3116)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO (cumple parámetros de potabilidad).\n• Bacteriológico: NO APTO (presencia de coliformes totales y fecales).\n• Conclusión: NO APTA para consumo humano. Requiere desinfección urgente de tanque de reserva y cañerías.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 1,80 Ω (Jabalina principal instalación general compartida, esquema TT con DD).\n• Conclusión: ADECUADO / CONFORME (adecuada disipación < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Conclusión: Valores relevados; se recomienda mantenimiento de tubos y redistribución para evitar deslumbramientos y sombras en puestos de atención.",
        "RUIDO: Relevado.\n• Conclusión: No se realizaron mediciones por no evidenciarse exposición sonora significativa.",
        "30712147829_ADMINISTRADORATRIBUTARIAER_CRESPO_38_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "16",
        "Etchevehere 375, Diamante, Entre Ríos (CP 3105)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: NO APTO por ARSÉNICO por encima del límite del Código Alimentario Argentino.\n• Bacteriológico: NO APTO (contaminación microbiológica significativa: coliformes totales elevados y coliformes termotolerantes/fecales).\n• Conclusión: NO APTA para consumo humano. Doble no conformidad (química y microbiológica). Riesgo sanitario.",
        "NO SE PUDO REALIZAR.\n• Conclusión: No fue posible identificar ni acceder a la toma de tierra (jabalina), sumado a tablero antiguo sin condiciones seguras de medición. Se recomendó adecuación eléctrica urgente y posterior medición.",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Conclusión: Relevada según anexo técnico. Se recomienda sustitución de artefactos defectuosos e instalación de luces de emergencia.",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica.",
        "30712147829_ADMINISTRADORATRIBUTARIAER_DIAMANTE_16_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "18",
        "Lavalle 265, Federal, Entre Ríos (CP 3180)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO (dentro de parámetros aceptados).\n• Bacteriológico: NO APTO (microorganismos aerobios mesófilos superiores a los admitidos).\n• Conclusión: NO APTA para consumo humano. Suministro proviene de tanque general del Centro Cívico; se requiere limpieza/desinfección institucional del tanque.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 1,14 Ω (Esquema TT con DD).\n• Conclusión: ADECUADO / CONFORME (condiciones seguras < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• 3 puntos: Mesa de entradas (08:30 hs), Oficina administrativa (08:30 hs: 295 lux vs 300 req), Archivo.\n• Conclusión: NO CUMPLE en sectores administrativos (valores por debajo del requerimiento de 300 lux); requiere refuerzo lumínico.",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica.",
        "30712147829_ADMINISTRADORATRIBUTARIAER_FEDERAL_18_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "71",
        "Sarmiento 121, Gualeguay, Entre Ríos (CP 2840)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO (valores compatibles con estándares normativos).\n• Bacteriológico: NO APTO (presencia de contaminación microbiológica).\n• Conclusión: NO APTA para consumo humano. Riesgo de enfermedades hídricas; se requiere limpieza/desinfección de tanques y nuevo control.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 6,10 Ω (Medición con telurímetro, esquema TT).\n• Conclusión: CUMPLE / CONFORME (inferior al límite de 40 Ω de Res. 900/15).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Puntos: Mesa de entradas (200 lux), Oficina trabajo general (Mixta, 500 lux req), Archivo.\n• Conclusión: NO CUMPLE (el valor medido en la Oficina es inferior al valor legal requerido de 500 Lux).",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica durante la visita.",
        "30712147829_ADMINISTRADORATRIBUTARIAER_GUALEGUAY_71_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "6",
        "Avenida Pancho Ramírez 71, Sauce de Luna, Entre Ríos (CP 3144)",
        "SÍ SE REALIZÓ.\n• Bacteriológico: NO APTO (presencia de contaminación microbiológica).\n• Fisicoquímico: NO APTO por RESIDUO A 105°C ELEVADO (1710 mg/l vs máx normativo 1500 mg/l).\n• Conclusión: NO APTA para consumo humano (riesgo de trastornos gastrointestinales).",
        "NO SE PUDO REALIZAR.\n• Conclusión: No se realizó medición de resistencia de puesta a tierra debido a la imposibilidad de acceso al terreno al momento del relevamiento. Se recomendó efectuarla ni bien se disponga de acceso.",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Puntos: Mesa de entradas, Oficina, Archivo.\n• Conclusión: NO CUMPLE (el valor medido en la Mesa de Entrada es inferior al valor requerido legalmente).",
        "RUIDO: Relevado.\n• Conclusión: No se realiza medición acústica instrumental dado que se perciben niveles muy bajos de ruido (riesgo bajo de contaminación sonora).",
        "30712147829_ADMINISTRADORATRIBUTARIAER_SAUCEDELUNA_6_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "9",
        "Hipólito Yrigoyen y Vandor, Bovril, Entre Ríos (CP 3142)",
        "SÍ SE REALIZÓ.\n• Bacteriológico: APTO (no se detecta contaminación microbiológica).\n• Fisicoquímico: NO APTO por ARSÉNICO ELEVADO (0,025 mg/l vs valor de referencia 0,01 mg/l).\n• Conclusión: NO APTA para consumo humano por presencia de arsénico por encima del Código Alimentario Argentino.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 0,26 Ω (Esquema TT, disyuntor presente).\n• Conclusión: ADECUADO / CUMPLE (dentro de parámetros reglamentarios < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• 3 puntos: Mesa de entrada (200 lux), Oficina (500 lux), Archivo (200 lux).\n• Conclusión: NO CUMPLE en sector Archivo (el valor medido en el Archivo es inferior al valor requerido legalmente).",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica durante la visita.",
        "30712147829_ADMINSITRADORATRIBUTARIAER_BOVRIL_9_ABRIL.2026_INFORME TECNICO.pdf"
    ],
    [
        "Abril",
        "30-71214782-9",
        "ADMINISTRADORA TRIBUTARIA DE LA PROVINCIA DE ENTRE RÍOS (ATER)",
        "49",
        "San Martín 388, General Ramírez, Entre Ríos (CP 3164)",
        "SÍ SE REALIZÓ.\n• Fisicoquímico: APTO (valores dentro de límites admitidos).\n• Bacteriológico: APTO (sin coliformes totales, E. coli ni patógenos).\n• Conclusión: APTA PARA CONSUMO HUMANO conforme al Art. 982 del Código Alimentario Argentino.",
        "SÍ SE REALIZÓ (Res. SRT 900/15).\n• Valor medido: 0,34 Ω (Jabalina interna patio, esquema TT con disyuntor DD).\n• Conclusión: ADECUADO / CONFORME (dentro de parámetros de seguridad < 40 Ω).",
        "SÍ SE REALIZÓ (Res. SRT 84/12).\n• Puntos: Mesa de entradas (300 lux req), Oficina administrativa (493 lux vs 300 req), Archivo (198 lux vs 200 req).\n• Conclusión: CUMPLE prácticamente en su totalidad (condiciones adecuadas de iluminación).",
        "RUIDO: Relevado durante la inspección.\n• Conclusión: No se detecta contaminación acústica.",
        "30712147829_ADMINSITRADORATRIBUTARIAER_RAMIREZ_49_ABRIL.2026_INFORME TECNICO.pdf"
    ]
]

# Write to Excel
# Header formatting
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ws.append(headers)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = header_align
ws.row_dimensions[1].height = 35

# Data formatting
data_font = Font(name="Calibri", size=10)
top_left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
center_top_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

for row_idx, r in enumerate(rows, start=2):
    ws.append(r)
    ws.row_dimensions[row_idx].height = 95
    for col_idx, val in enumerate(r, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = data_font
        cell.border = thin_border
        if col_idx in [1, 2, 4]:
            cell.alignment = center_top_align
        else:
            cell.alignment = top_left_align

# Adjust column widths
col_widths = {
    1: 14, # Mes
    2: 16, # CUIT
    3: 32, # RS
    4: 10, # Est
    5: 35, # Domicilio
    6: 45, # Agua
    7: 35, # PAT
    8: 40, # Iluminacion
    9: 38, # Otras
    10: 45 # Archivo
}

for col_idx, width in col_widths.items():
    col_letter = get_column_letter(col_idx)
    ws.column_dimensions[col_letter].width = width

excel_path = r'c:\Users\cript\OneDrive\Escritorio\Antigravity\Informes IAPSER\Relevamiento_Mediciones_Abril_2026.xlsx'
wb.save(excel_path)
print(f"Excel file successfully generated at: {excel_path}")
